"""발표 초안 report/draft.pptx 생성 (python-pptx) — 그래프는 06_export_figures.py가 만든 png 사용."""
# %% [7-0] 설정
import re
import sys
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / 'outputs' / 'figures'
OUT = ROOT / 'report' / 'draft.pptx'

FONT = '맑은 고딕'
GREEN = RGBColor(0x1E, 0x9E, 0x5A)
TEXT = RGBColor(0x33, 0x33, 0x33)
SUB = RGBColor(0x88, 0x88, 0x88)
BOX = RGBColor(0xE8, 0xF5, 0xEE)
CARD = RGBColor(0xF5, 0xF6, 0xF5)
BORDER = RGBColor(0xCC, 0xCC, 0xCC)
W, H = 13.333, 7.5
LEFT = 0.6
CONTENT_W = W - 2 * LEFT
TITLE_FONT = ImageFont.truetype(r'C:\Windows\Fonts\malgunbd.ttf', 280)  # 28pt × 10


# %% [7-1] 공통 도구: 글꼴·문단·제목·그림·박스
def style_run(run, size, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rpr.set('lang', 'ko-KR')
    rpr.set('altLang', 'en-US')
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', FONT)


def add_rich(paragraph, text, size, color=TEXT, bold=False):
    """**굵게** 표시한 부분은 초록 굵게."""
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        for j, piece in enumerate(part.split('\n')):
            if j > 0:
                paragraph.add_line_break()
            if not piece:
                continue
            run = paragraph.add_run()
            run.text = piece
            style_run(run, size, bold=bold or i % 2 == 1, color=GREEN if i % 2 == 1 else color)


def set_bullet(paragraph, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set('marL', str(Emu(Inches(0.3 + 0.32 * level))))
    ppr.set('indent', str(-Emu(Inches(0.24))))
    bu_clr = ppr.makeelement(qn('a:buClr'), {})
    clr = bu_clr.makeelement(qn('a:srgbClr'), {'val': '1E9E5A' if level == 0 else '888888'})
    bu_clr.append(clr)
    bu_font = ppr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    bu_char = ppr.makeelement(qn('a:buChar'), {'char': '•' if level == 0 else '–'})
    for el in (bu_clr, bu_font, bu_char):
        ppr.append(el)


def text_box(slide, x, y, w, h, items, size=17, bullets=True, color=TEXT, space_after=10, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, bold=False, line_spacing=1.15):
    """items: 문자열 또는 (문자열, 수준). 수준 1은 작은 보조 글머리표."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', Inches(0.05))
    for i, item in enumerate(items):
        text, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p._p.get_or_add_pPr().set('eaLnBrk', '0')  # 한글 단어 중간에서 줄바꿈하지 않음
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        sz = size if level == 0 else size - 2
        add_rich(p, text, sz, color=color if level == 0 else SUB, bold=bold)
        if bullets:
            set_bullet(p, level)
    return tb


def title_lines(title):
    if '\n' in title:
        return title.count('\n') + 1
    return 1 if TITLE_FONT.getlength(title) / 10 / 72 <= CONTENT_W - 0.3 else 2


def add_title(slide, label, title):
    """섹션 라벨 + 제목 + 짧은 초록 선. 본문 시작 y를 돌려줌."""
    text_box(slide, LEFT, 0.32, CONTENT_W, 0.3, [label], size=11, bullets=False, color=GREEN, bold=True)
    lines = title_lines(title)
    title_h = 0.55 * lines
    text_box(slide, LEFT, 0.6, CONTENT_W, title_h, [title], size=28, bullets=False, bold=True, space_after=0,
             line_spacing=1.0)
    line_y = 0.6 + title_h + 0.12
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(LEFT + 0.05), Inches(line_y), Inches(1.5), Pt(3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()
    return line_y + 0.3


def add_picture(slide, name, x, y, w, h, align='center'):
    """박스 안에 비율 유지해 최대 크기로, 얇은 회색 테두리."""
    path = FIG / f'{name}.png'
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == 'center' else (x if align == 'left' else x + w - pw)
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = BORDER
    pic.line.width = Pt(0.75)
    return pic


def rounded_box(slide, x, y, w, h, fill, radius=0.12):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def bottom_box(slide, text, y=6.5, h=0.62):
    rounded_box(slide, LEFT, y, CONTENT_W, h, BOX, radius=0.25)
    text_box(slide, LEFT + 0.25, y, CONTENT_W - 0.5, h, [text], size=17, bullets=False, space_after=0,
             anchor=MSO_ANCHOR.MIDDLE)


def page_number(slide, n):
    text_box(slide, W - 1.2, H - 0.45, 0.8, 0.3, [str(n)], size=10, bullets=False, color=SUB,
             align=PP_ALIGN.RIGHT, space_after=0)


def stat(slide, x, y, w, number, label, num_size=30):
    text_box(slide, x, y, w, 0.6, [f'**{number}**'], size=num_size, bullets=False, space_after=0, line_spacing=1.0)
    text_box(slide, x, y + 0.62, w, 0.45, [label], size=13, bullets=False, color=SUB, space_after=0)


def card(slide, x, y, w, h, header, items, size=16, header_size=18):
    rounded_box(slide, x, y, w, h, CARD, radius=0.08)
    n_lines = header.count('\n') + 1
    head_h = header_size / 72 * 1.25 * n_lines
    text_box(slide, x + 0.2, y + 0.15, w - 0.4, head_h + 0.1, [header], size=header_size, bullets=False,
             color=GREEN, bold=True, space_after=0, line_spacing=1.05)
    if items:
        text_box(slide, x + 0.2, y + 0.3 + head_h, w - 0.4, h - head_h - 0.4, items, size=size, space_after=6)


def simple_table(slide, x, y, w, rows, col_ws, size=14, head=None, row_h=0.5):
    """첫 칸은 초록 굵게(연한 회색 바탕), 나머지는 본문. head가 있으면 맨 위 머리글 줄."""
    data = ([head] if head else []) + rows
    shape = slide.shapes.add_table(len(data), len(col_ws), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(data)))
    tbl = shape.table
    tbl.first_row = bool(head)
    tbl.horz_banding = False
    for j, cw_ in enumerate(col_ws):
        tbl.columns[j].width = Inches(cw_)
    for i, row in enumerate(data):
        tbl.rows[i].height = Inches(row_h)
        for j, text in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            is_head = head and i == 0
            cell.fill.fore_color.rgb = BOX if is_head else (CARD if j == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para._p.get_or_add_pPr().set('eaLnBrk', '0')
            add_rich(para, text, size, color=GREEN if (j == 0 or is_head) else TEXT, bold=(j == 0 or is_head))
    return shape


def notes(slide, lines):
    slide.notes_slide.notes_text_frame.text = chr(10).join(lines)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
page = 0


def new_slide():
    global page
    page += 1
    s = prs.slides.add_slide(BLANK)
    if page > 1:
        page_number(s, page)
    return s


# %% [7-2] 1. 표지
s = new_slide()
text_box(s, 0.8, 2.15, W - 1.6, 0.9, ['따릉이는 지하철의 경쟁자일까, 동반자일까?'], size=40, bullets=False, bold=True,
         align=PP_ALIGN.CENTER, space_after=0)
text_box(s, 0.8, 3.15, W - 1.6, 0.5, ['2019 vs 2026 (5~6월) · 서울 지하철 승하차 × 따릉이 대여이력'], size=20,
         bullets=False, color=GREEN, bold=True, align=PP_ALIGN.CENTER, space_after=0)
text_box(s, 0.8, 4.35, W - 1.6, 0.4, ['202310885_하동민'], size=16, bullets=False, color=SUB,
         align=PP_ALIGN.CENTER, space_after=0)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(H - 0.55), Inches(W), Inches(0.55))
band.fill.solid()
band.fill.fore_color.rgb = GREEN
band.line.fill.background()

# %% [7-3] 2. 개요
s = new_slide()
top = add_title(s, 'Chapter 1 · 개요', 'Chapter 1. 개요')
cw, ch, gap = (CONTENT_W - 0.3) / 2, 2.3, 0.3
card(s, LEFT, top, cw, ch, '동기', [
    '2019 서울 지하철 승하차 EDA의 후속 분석',
    '7년 사이 따릉이 이용은 **거의 2배**',
    '늘어난 따릉이, 지하철 승객을 가져갔을까?'])
card(s, LEFT + cw + gap, top, cw, ch, '확인할 것', [
    '**가설①(경쟁)** 역간 단거리 따릉이 이동이 많이 는 역일수록 지하철 이용 증가율이 더 낮았나?',
    '**가설②(동반)** 출근 시간 역 근처 반납이 몰리고, 역 연계 이동이 는 역일수록 지하철 증가율도 더 높았나?'], size=15)
card(s, LEFT, top + ch + gap, cw, ch, '최종 흐름', [
    '큰 그림 → 가설①(경쟁) → 가설②(동반)',
    '→ 추가 확인(견고성) → 인사이트 → 결론'])
card(s, LEFT + cw + gap, top + ch + gap, cw, ch, '가져온 외부 데이터', [
    '서울시 **따릉이 대여이력**·대여소 위치 (5~6월)',
    '이상 기록 제외 후 **390만 / 719만 건** (원본의 86~88%)',
    '대여소별 월별 이용 (분석 기간 선정)'], size=15)

# %% [7-4] 3. 큰 그림
s = new_slide()
top = add_title(s, '큰 그림', '7년 사이 따릉이는 어떻게 변했나?')
sw = CONTENT_W / 4
for i, (num, label) in enumerate([('+88%', '평일 하루 이용 (6.7만→12.6만)'),
                                  ('1.8배', '이용된 대여소 (1,512 → 2,777곳)'),
                                  ('2.0배', '역 연계 이동 (평일 하루)'),
                                  ('2.7배', '7~9시 반납 (하루 전체는 1.9배)')]):
    stat(s, LEFT + i * sw, top - 0.05, sw - 0.2, num, label)
gy = top + 1.15
add_picture(s, '02_overview', LEFT, gy, CONTENT_W / 2 - 0.15, 6.35 - gy)
add_picture(s, '00_bike_monthly', LEFT + CONTENT_W / 2 + 0.15, gy, CONTENT_W / 2 - 0.15, 6.35 - gy)
bottom_box(s, '지하철 -2.9%일 때 따릉이 +88% · 4월/5월 0.67배 → 1.01배: 레저용에서 **일상 교통수단**으로')

# %% [7-5] 4. 데이터와 방법
s = new_slide()
top = add_title(s, '데이터 · 방법', '데이터와 방법')
text_box(s, LEFT, top, 5.7, 6.9 - top, [
    '기간: 2019·2026 **5~6월** (공휴일·선거일 제외)',
    ('평일 40일 / 39일, 주말 18일 / 18일', 1),
    '지하철: 두 해 모두 있는 서울 역 **295곳**',
    ('역 이용량 = 승차 + 하차', 1),
    '따릉이: 이상 기록 제외 후 **390만 / 719만 건** (원본의 86~88%)',
    '대여소 → 반경 **500m** 안 가장 가까운 역에 배정',
    '이동 3유형',
    ('역간 단거리: 서로 다른 두 역 근처, 2.5km 이내', 1),
    ('역 연계: 한쪽 끝만 역 근처', 1),
    ('기타: 그 밖의 이동', 1),
    '증가율 = 2026 ÷ 2019 **평일 하루 평균**'], size=17, space_after=7)
add_picture(s, '01_map_seoul', LEFT + 5.9, top - 0.05, CONTENT_W - 5.9, 6.95 - top)

# %% [7-5b] 4-1. 분석 기준 한눈에
s = new_slide()
top = add_title(s, '데이터 · 방법', '분석 기준 한눈에')
simple_table(s, LEFT, top, CONTENT_W, [
    ['데이터', '지하철 역별 승하차 · 따릉이 대여이력 (2019·2026 5~6월) · 대여소 위치 (21.1월·26.6월판, 2019년판 없음) · 개인정보 미사용'],
    ['기간·날짜', '5~6월 (두 해 모두 성수기) · **평일 중심** 40 / 39일, 주말 18 / 18일 · 공휴일·근로자의날·선거일 제외 · 출근 = 07:00~08:59 반납'],
    ['역 선정', '527 → **295곳**: 1km 안 대여소 있는 서울 역 · 경기 5곳 제외 · 두 해 모두 있는 역 · 총신대입구 = 이수 · 양평 제외 · 신림선·서해선 승객 제외'],
    ['이상 기록 제외', '반납 없음 · 2분 미만 · 2시간 초과 (실제 시각 차이) · 같은 대여소 반납 · 위치 모름 제외 → 원본의 **86% / 88%** 남음'],
    ['주요 기준', '대여소 → **500m** 안 가장 가까운 역 1곳 · 역간 단거리 = 두 역 **2.5km** 이내 (1.5·3km도 확인) · 2019 평일 10건 미만 역 제외 (**282 / 270곳**)'],
], col_ws=[1.9, CONTENT_W - 1.9], size=15, row_h=0.95)

# %% [7-6] 5. 가설① 설정
Q1 = '역간 단거리 따릉이 이동이 많이 는 역일수록\n지하철 이용 증가율이 더 낮았나?'
s = new_slide()
top = add_title(s, '가설① · 설정', Q1)
text_box(s, LEFT, top, 5.6, 6.35 - top, [
    '**H0** 역간 단거리 이동 증가율과 지하철 증가율 사이에 상관이 없다 (ρ=0)',
    '**H1** 상관이 있다 (ρ≠0)',
    '①-1 Spearman · 역 **n=282**',
    '①-2 순위합 검정(W) · 증가율 중앙값 기준',
    ('두 그룹: 많이 는 역 141 / 적게 는 역 141', 1),
    '대상: 2019 평일 하루 10건 이상 역 (양측 검정)'], size=16, space_after=9)
add_picture(s, '11_hyp1_dist', LEFT + 5.8, top, CONTENT_W - 5.8, 6.35 - top)
bottom_box(s, '두 증가율 모두 정규 아님 (p<0.05) → 순위 기반 **Spearman · 순위합 검정(W)**')

# %% [7-7] 6. 가설① 결과
s = new_slide()
top = add_title(s, '가설① · 결과', Q1)
gh = 5.35 - top
add_picture(s, '04_hyp1_scatter', LEFT, top, CONTENT_W * 0.56, gh)
add_picture(s, '05_hyp1_bar', LEFT + CONTENT_W * 0.58, top, CONTENT_W * 0.42, gh)
text_box(s, LEFT, 5.45, CONTENT_W, 1.0, [
    '①-1 Spearman **ρ = 0.24** (양, 약함), p<0.05 · 역 n=282 · 거리 기준 1.5 / 3km에서도 양(+), p<0.05',
    '①-2 순위합 검정(W) 지하철 증가율: 많이 는 역 **+0.2%** vs 적게 는 역 **-5.1%**, p<0.05'], size=16,
    space_after=4)
bottom_box(s, '음(-)이 아니라 양(+) → **경쟁 근거 없음**   다음 질문: 그렇다면 따릉이는 지하철의 동반자일까?')

# %% [7-8] 7. 가설② 설정
s = new_slide()
top = add_title(s, '가설② · 설정', '따릉이는 지하철과 이어지는 동반자가 됐나?')
cw, ch = 5.75, (6.3 - top - 0.2) / 2
card(s, LEFT, top, cw, ch, '②-1 평일 출근 시간(7~9시)에\n역 근처 반납이 더 몰리게 됐나?', [
    '**H0** 두 해의 7~9시 역 근처 반납 비중 분포가 같다',
    '**H1** 분포가 다르다 · 순위합 검정(W)',
    '단위: 평일 하루 **n=40 / 39**'], size=15, header_size=16)
card(s, LEFT, top + ch + 0.2, cw, ch, '②-2 역 연계 이동이 많이 는 역일수록\n지하철 증가율도 더 높았나?', [
    '**H0** 역 연계 이동 증가율과 지하철 증가율은 상관이 없다 (ρ=0)',
    '**H1** 상관이 있다 (ρ≠0) · Spearman',
    '단위: 역 **n=270**'], size=15, header_size=16)
add_picture(s, '12_hyp2_dist', LEFT + cw + 0.2, top, CONTENT_W - cw - 0.2, 6.3 - top)
bottom_box(s, '모두 정규 아님 (p<0.05) → 순위 기반 **순위합 검정(W) · Spearman**')

# %% [7-9] 8. 가설② 결과 ②-1
s = new_slide()
top = add_title(s, '가설② · 결과 ②-1', '평일 출근 시간(7~9시)에 역 근처 반납이 더 몰리게 됐나?')
gh = 5.35 - top
add_picture(s, '06_hyp2_hourly', LEFT, top, CONTENT_W * 0.6, gh)
add_picture(s, '08_hyp2_share_bar', LEFT + CONTENT_W * 0.62, top, CONTENT_W * 0.38, gh)
text_box(s, LEFT, 5.45, CONTENT_W, 1.0, [
    '7~9시 역 근처 반납 비중 **11.5% → 16.4%**, 하루 **4,922 → 13,190건** · 순위합 검정(W), p<0.05 (평일 n=40 / 39)',
    '단, 역에서 먼 반납도 7~9시 비중 12.2% → 16.3%로 같이 증가'], size=16, space_after=4)
bottom_box(s, '따릉이는 **출근 수단**으로 커졌다 — 역 근처만이 아니라 따릉이 전체에서')

# %% [7-10] 9. 가설② 결과 ②-2
s = new_slide()
top = add_title(s, '가설② · 결과 ②-2', '역 연계 이동이 많이 는 역일수록\n지하철 증가율도 더 높았나?')
stat(s, LEFT, top + 0.1, 2.7, 'ρ = 0.28', '②-2 역 연계 이동', num_size=32)
stat(s, LEFT + 2.9, top + 0.1, 2.7, 'ρ = 0.24', '①-1 역간 단거리 이동 (참고)', num_size=32)
text_box(s, LEFT, top + 1.45, 5.5, 6.35 - top - 1.45, [
    'Spearman **ρ = 0.28** (양, 약함), p<0.05',
    ('역 n=270', 1),
    '결과 읽는 법으로는 "동반 지지"',
    '그런데 역간 단거리 이동(①-1)도 비슷한 크기의 양(+)',
    '→ 연결 방식과 상관없이, 따릉이가 많이 는 동네는 지하철도 덜 줄었다'], size=17, space_after=10)
add_picture(s, '09_hyp2_scatter', LEFT + 5.7, top, CONTENT_W - 5.7, 6.35 - top)
bottom_box(s, '다음 질문: 그렇다면 왜 **둘 다 양(+)**일까?')

# %% [7-11] 10. 추가 확인: 견고성
s = new_slide()
top = add_title(s, '추가 확인 · 견고성', '같이 늘어난 건 따릉이 때문일까?')
text_box(s, LEFT, top, 4.9, 6.35 - top, [
    '조건을 바꿔 같은 검정을 다시 했다',
    ('(a) 두 해 모두 있던 대여소만', 1),
    ('(b) 하루 평균 대신 중앙값', 1),
    ('(c) 지하철 증가율 상위 5% 역 제외', 1),
    '(b)·(c): 양(+), p<0.05 **유지**',
    '(a): ρ **0.04 / 0.05**, p>0.05로 **사라짐**',
    ('단, (a)는 2026 이동이 28%만 남음', 1),
    '출근 시간 증가(②-1)는 (a)에서도 유지'], size=17, space_after=8)
add_picture(s, '10_robustness_rho', LEFT + 5.0, top, CONTENT_W - 5.0, 6.35 - top)
bottom_box(s, '같이 늘어난 건 **성장 지역의 새 대여소** 효과 (추정) · 출근 시간 증가는 그대로')

# %% [7-12] 11. 인사이트
s = new_slide()
top = add_title(s, '인사이트', '인사이트: 어디에 무엇을 할까')
insights = [('입주·개발 지역 대여소 선제 설치', '새 대여소가 생긴 지역에서 수요 증가 (추가 확인) · Top 10 중 6개가 마곡·발산'),
            ('출근 시간 반납 거치대 확충·재배치', '평일 7~9시 반납 2.7배'),
            ('역간 이동 집중 구간 우선 관리', '마곡↔마곡나루 평일 하루 803건(7년 새 약 7.5배) → 양쪽 역 대여소 확충·재배치, 자전거 길 점검'),
            ('운영 지표를 평일 출퇴근 중심으로', '4월/5월 0.67배 → 1.01배, 계절 편차 감소')]
row_h = (6.75 - top) / 4
for i, (head, sub) in enumerate(insights):
    y = top + i * row_h
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(LEFT), Inches(y + 0.08), Inches(0.5), Inches(0.5))
    circ.fill.solid()
    circ.fill.fore_color.rgb = GREEN
    circ.line.fill.background()
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = str(i + 1)
    style_run(r, 18, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    text_box(s, LEFT + 0.7, y + 0.02, 5.1, 0.45, [head], size=18, bullets=False, bold=True, space_after=0)
    text_box(s, LEFT + 0.7, y + 0.48, 5.1, row_h - 0.5, [sub], size=14, bullets=False, color=SUB, space_after=0,
             line_spacing=1.05)
add_picture(s, '07_top10', LEFT + 5.9, top, CONTENT_W - 5.9, 6.75 - top)

# %% [7-13] 12. 한계
s = new_slide()
top = add_title(s, '한계', '한계')
limits = [('상관이지 인과가 아니다', '"따릉이 때문에"라고 말할 수 없음'),
          ('환승 여부를 모른다', '같은 사람이 갈아탔는지 알 수 없음 → 동네 단위 경향'),
          ('②-1 날짜 독립성', '평일 하루들은 날씨 등 공통 영향을 받음'),
          ('(a) 표본 감소', '공통 대여소만 남기면 2026 이동 28%'),
          ('개발 지역 영향', '입주·재건축 지역을 직접 통제하지 못함'),
          ('이용 건수 기준', '이용자 수가 아닌 이용 건수')]
cw, ch = (CONTENT_W - 0.6) / 3, (6.85 - top - 0.3) / 2
for i, (head, sub) in enumerate(limits):
    x = LEFT + (i % 3) * (cw + 0.3)
    y = top + (i // 3) * (ch + 0.3)
    rounded_box(s, x, y, cw, ch, CARD, radius=0.08)
    text_box(s, x + 0.25, y + 0.25, 0.8, 0.6, [f'**{i + 1:02d}**'], size=26, bullets=False, space_after=0)
    text_box(s, x + 0.25, y + 0.95, cw - 0.5, 0.5, [head], size=18, bullets=False, bold=True, space_after=0)
    text_box(s, x + 0.25, y + 1.45, cw - 0.5, ch - 1.55, [sub], size=15, bullets=False, color=SUB, space_after=0)

# %% [7-14] 13. 결론
s = new_slide()
top = add_title(s, '결론', '결론')
conclusions = [('따릉이는 지하철의\n**경쟁자가 아니다**', '어떤 조건에서도 유의한 음(-)의 관계 없음'),
               ('같이 늘어난 건\n**대여소 신설 효과**', '두 해 공통 대여소만 보면 사라짐 (2026 이동 28%만 남는 한계)'),
               ('따릉이는\n**출근 수단**이 됐다', '7~9시 역 근처 반납 비중 11.5% → 16.4%, 검정한 모든 조건에서 유지')]
cw, ch = (CONTENT_W - 0.6) / 3, 6.3 - top - 0.25
for i, (head, sub) in enumerate(conclusions):
    x = LEFT + i * (cw + 0.3)
    rounded_box(s, x, top, cw, ch, CARD, radius=0.06)
    text_box(s, x + 0.3, top + 0.3, 0.8, 0.6, [f'**{i + 1}**'], size=30, bullets=False, space_after=0)
    tb = text_box(s, x + 0.3, top + 1.0, cw - 0.6, 1.4, [head], size=24, bullets=False,
                  bold=True, space_after=0)
    text_box(s, x + 0.3, top + 2.55, cw - 0.6, ch - 2.7, [sub], size=16, bullets=False, color=SUB, space_after=0)
bottom_box(s, '→ 따릉이는 지하철을 뺏지도 끌어올리지도 않는, **출퇴근의 빈틈을 채우는 일상 교통수단**', y=6.45)

# %% [7-15] 부록. 예상 질문 Q&A
s = new_slide()
top = add_title(s, '부록', '예상 질문 Q&A')
simple_table(s, LEFT, top - 0.1, CONTENT_W, [
    ['왜 반경 500m?', '역 출구에서 걸어갈 만한 거리 (도보 약 5~7분, 추정). 다른 반경은 확인 안 함'],
    ['2.5km는 임의 아닌가?', '맞다. 1.5km·3km로 바꿔도 양(+), p<0.05로 같았다'],
    ['따릉이 때문에 지하철이 늘었나?', '아니다. 상관일 뿐이고, 공통 대여소만 보면 관계가 사라진다'],
    ['같은 사람이 갈아탔나?', '확인 못 함. 개인 단위로 이어진 데이터가 아니다 → 동네 단위 경향'],
    ['이용자 수인가?', '아니다. 이용 건수다'],
    ['왜 5~6월?', '두 해 모두 성수기. 2019년 4월은 5월의 2/3 수준이었다'],
    ['같은 대여소 반납은 왜 뺐나?', '출발·도착이 같아 이동 방향을 알 수 없다 (왕복 산책 등 추정)'],
    ['왜 평일 중심?', '평일이 지하철 이용의 중심이고, 따릉이가 지하철을 대신하거나 이어주는 건 일상 이동에서 생김. 주말은 레저 이동이 섞임 (출근 시간만 본 건 ②-1)'],
    ['대여소가 2배로 늘었는데 공정한가?', '공통 대여소만으로도 확인했다. 출근 시간 증가는 그대로 유지'],
    ['날씨는 반영했나?', '직접은 안 했다. 하루 값 중앙값으로 바꿔도 결과는 같았다'],
    ['에러바가 겹치는데 왜 유의?', '에러바는 중앙값 하나의 범위, 순위합 검정은 두 그룹 전체 순위를 비교'],
    ['검정을 여러 번 해서 우연 아닌가?', '핵심 검정 4개로 미리 정함. 4개를 고려한 엄격한 기준(0.05÷4=0.0125)에서도 모두 유의'],
], col_ws=[3.7, CONTENT_W - 3.7], size=13, head=['질문', '한 줄 답'], row_h=0.41)

# %% [7-16] 발표자 노트: 슬라이드별 나올 만한 질문·답·관련 기준
SLIDE_NOTES = [
    ['[표지] 한 줄 소개: 7년 사이 2배가 된 따릉이가 지하철 승객을 가져갔는지, 아니면 함께 가는지 본 분석',
     'Q. 왜 2019 vs 2026? → 기존 2019 지하철 EDA를 이어서, 같은 계절(5~6월) 7년 차이를 비교'],
    ['Q. 외부 데이터가 뭔가? → 서울 열린데이터광장 따릉이 대여이력 + 대여소 위치 + 대여소별 월별 이용',
     'Q. 개인정보는? → 생년·성별·이용자 종류 칸은 읽지도 쓰지도 않았다',
     '기준: 가설①(경쟁)·가설②(동반) 각각 검정 2개씩, 총 4개'],
    ['Q. 지하철 숫자는 어디 기준? → 분석 대상 295개 역 승차+하차 합계, 평일 하루 평균 1,304.9만 → 1,267.1만 명 (-2.9%)',
     'Q. 따릉이 숫자는? → 이상 기록 제외 후 서울 전체 이동, 평일 하루 66,945 → 125,878건 (+88%)',
     'Q. "일상 교통수단"의 근거? → 4월/5월 이용 비율 0.67배 → 1.01배 (계절 편차 감소), 7~9시 반납 2.7배 (하루 전체 1.9배)',
     '에러바는 날짜 단위 95% 신뢰구간 (평일 n=40/39일)'],
    ['Q. 왜 500m? → 역 출구에서 걸어갈 만한 거리 (도보 약 5~7분, 추정). 다른 반경은 확인 안 함',
     'Q. 대여소가 두 역 근처에 있으면? → 가장 가까운 역 1곳에만 배정해 중복 집계 없음',
     'Q. 서울 역은 어떻게 골랐나? → 반경 1km 안에 따릉이 대여소가 있는 역 + 경기 5곳(역곡·철산·석수·갈매·지축) 직접 제외',
     'Q. 2019년 대여소 위치는? → 2019년판이 없어 26.6월판 우선, 없으면 21.1월판. 둘 다 없는 대여소 이용은 제외'],
    ['Q. 527곳에서 295곳으로 어떻게 줄었나? → 양평 1, 좌표 매칭 안 됨(2026 신규역) 18, 경기 5, 서울 밖 208곳 제외',
     'Q. 근로자의날은? → 두 해 기준을 맞추려고 둘 다 제외',
     'Q. 이용시간은 어떻게 계산? → 원본 칸이 아니라 실제 시각 차이 (2019 칸 불일치 2.25%)',
     'Q. 새 노선은? → 신림선·서해선 승객 행만 2026에서 뺐다 (보라매 39.0%, 샛강 32.5% 등)'],
    ['Q. 왜 Spearman·순위합 검정(W)? → 두 증가율 모두 정규분포 아님 (Shapiro p<0.05)',
     'Q. 왜 10건 미만 역을 뺐나? → 2019 값이 작으면 증가율이 수천 %로 튀어서. 남은 역 282곳',
     'Q. 2.5km 기준은? → 약 1~2정거장. 1.5km·3km도 확인함',
     '양측검정, 유의수준 0.05'],
    ['Q. ρ=0.24는 얼마나 강한가? → 약한 양의 관계. 경쟁이라면 음(-)이어야 한다',
     'Q. 거리 기준 바꾸면? → 1.5km ρ=0.21, 3km ρ=0.24, 모두 p<0.05 유지',
     'Q. 튀는 점(둔촌오륜·마곡)은? → 재건축·신도시 지역 추정. 상위 5% 역을 빼도 ρ=0.17, p<0.05',
     '그래서 뭐? → 경쟁 근거 없음 → 다음 질문: 동반자인가?'],
    ['Q. 출근 시간은 언제? → 07:00~08:59 반납 (반납 시각 기준)',
     'Q. ②-1 단위가 왜 하루인가? → 날마다 역 근처 반납 중 7~9시 비중을 구해 2019 평일 40일 vs 2026 평일 39일 비교',
     'Q. 역 연계 이동 정의? → 출발·도착 중 한쪽만 역 500m 안',
     '모두 정규분포 아님 (Shapiro p<0.05) → 비모수 검정'],
    ['Q. 비중이 늘었나 건수가 늘었나? → 둘 다. 비중 11.5% → 16.4%, 하루 4,922 → 13,190건',
     'Q. 역 근처만 늘었나? → 아니다. 역에서 먼 반납도 7~9시 비중 12.2% → 16.3%로 같이 증가',
     'Q. 날씨 영향은? → 직접 반영 안 함. 2019년 7~9시 비중이 1~3%인 날 3일 (비 온 날 추정)',
     '한계: 하루들은 날씨 등 공통 영향을 받아 완전히 독립은 아님'],
    ['Q. ρ=0.28이면 동반 아닌가? → 읽는 법상 "동반 지지"지만, 역간 단거리(①-1)도 0.24로 비슷한 양(+)',
     'Q. 그게 왜 문제? → 연결 방식과 상관없이 같이 늘었다는 뜻 → 동네 성장 효과일 가능성',
     '대상 역 270곳 (2019 평일 역 연계 10건 이상)',
     '그래서 뭐? → 다음 장: 같이 늘어난 게 따릉이 때문인지 확인'],
    ['Q. (a)는 뭘 한 건가? → 2019·2026 모두 이용된 대여소 1,374곳 사이 이동만 남겨 대여소 증설 효과를 뺐다',
     'Q. 표본이 줄어서 사라진 것 아닌가? → 그럴 수 있다. 2026 이동이 27.8%만 남아 구분 못 함 (한계)',
     'Q. (b)(c)는? → 하루 값 중앙값, 지하철 급등 상위 5% 역 제외 → 모두 양(+), p<0.05 유지',
     '출근 시간 증가(②-1)는 (a)에서도 +4.6%p, p<0.05 유지 → 이용 방식 자체가 바뀜'],
    ['Q. 마곡·발산이 왜 많나? → 개발·입주 지역이라 역 사이 짧은 이동이 몰림 (추정)',
     'Q. 정책 근거는? → 1) 새 대여소가 생긴 지역에서 수요 증가(견고성 a) 2) 7~9시 반납 2.7배 3) 마곡↔마곡나루 평일 하루 803건(7년 새 약 7.5배)',
     'Q. 3번 "우선 관리"는 뭘 하나? → 양쪽 역 대여소 확충·재배치, 자전거 길 점검',
     'Q. 경쟁이 아니면 역간 이동을 늘려도 되나? → 지하철을 뺏는 증거가 없어 짧은 구간 보완 수단으로 볼 수 있다'],
    ['Q. 가장 큰 한계? → 상관이지 인과가 아니고, 같은 사람의 환승을 확인하지 못한 것',
     'Q. 다음에 할 것? → 동네 성장(인구·입주) 통제, 이용자 단위 데이터, 기상청 날씨 데이터 결합',
     '숫자는 이용자 수가 아니라 이용 건수'],
    ['결론 1: 어떤 조건에서도 유의한 음(-)의 관계 없음 (거리 3가지, 견고성 3가지 포함)',
     '결론 2: 공통 대여소만 보면 ①-1·②-2 관계가 사라짐 (단, 2026 이동 28%만 남는 한계)',
     '결론 3: 7~9시 비중 증가는 검정한 모든 조건(본 분석, 공통 대여소)에서 유지',
     'Q. 그래서 따릉이는 뭔가? → 지하철을 뺏지도 끌어올리지도 않는, 출퇴근의 빈틈을 채우는 일상 교통수단'],
    ['더 자세한 기준과 답은 docs/분석기준_QA.md, notebooks/walkthrough.ipynb 10장 참고',
     'Q. 왜 평일 중심? → 평일이 지하철 이용의 중심이고, 따릉이가 지하철을 대신하거나 이어주는 건 일상 이동에서 생김. 주말은 레저 이동이 섞임 (출근 시간만 본 건 ②-1)',
     'Q. 에러바가 겹치는데 왜 유의? → 에러바는 중앙값 하나의 불확실성 범위, 순위합 검정(W)은 두 그룹 전체의 순위를 비교한다',
     'Q. 여러 번 검정해 우연 아닌가? → 핵심 검정 4개로 미리 정함. 0.05÷4=0.0125 기준에서도 4개 모두 유의',
     'Q. 지하철은 왜 2.9% 줄었나? → 원인은 분석하지 않았다 (확인 안 함)'],
]
assert len(SLIDE_NOTES) == len(prs.slides), (len(SLIDE_NOTES), len(prs.slides))
for slide, lines in zip(prs.slides, SLIDE_NOTES):
    notes(slide, lines)

prs.save(OUT)
print(f'저장: {OUT} ({page}장)')
